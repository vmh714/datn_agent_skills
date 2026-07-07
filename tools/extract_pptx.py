from pptx import Presentation

try:
    prs = Presentation(r"d:\datn\report\Do_an_tot_nghiep_Vu_Manh_Hung\DATN_VuManhHung_slides.pptx")
    for i, slide in enumerate(prs.slides):
        title = ""
        if slide.shapes.title:
            title = slide.shapes.title.text.strip()
        print(f"\n--- Slide {i+1}: {title} ---")
        for j, shape in enumerate(slide.shapes):
            if shape.is_placeholder:
                continue
            try:
                l = shape.left.inches
                t = shape.top.inches
                w = shape.width.inches
                h = shape.height.inches
            except Exception:
                continue
            
            content_preview = ""
            if shape.has_text_frame:
                content_preview = "TEXT: " + shape.text[:50].replace("\n", " ")
            elif getattr(shape, "shape_type", None) == 13:
                content_preview = "PICTURE"
            elif shape.has_table:
                content_preview = "TABLE"
                
            print(f"left={l:.2f}, top={t:.2f}, width={w:.2f}, height={h:.2f} | {content_preview}")
except Exception as e:
    print("Error:", e)
