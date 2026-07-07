# -*- coding: utf-8 -*-
"""
Build slide bảo vệ DATN (Fall Detection Eldercare IoT + TinyML) — Vũ Mạnh Hưng.

>>> CONTENT-DRIVEN: nội dung slide nằm trong `slides_content.md` (SỬA FILE ĐÓ).
    Script này chỉ là RENDERER: đọc Markdown → dựng .pptx bằng python-pptx.
    Bản imperative cũ được backup ở `build_slides_legacy_imperative.py.bak`.

Số liệu CHUẨN từ báo cáo Chương 5 (17 ms @160MHz · 11,2 ms @240MHz · ~180× TCN).
Rule: visual-first — MỌI slide có hình/sơ đồ/bảng; ảnh trên, chữ dưới (hoặc trong notes).
Chạy: python build_slides.py
"""
import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu, Length
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import PP_PLACEHOLDER, MSO_SHAPE
from pptx.oxml.ns import qn

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # d:\datn
REPORT = os.path.join(BASE, "report", "Do_an_tot_nghiep_Vu_Manh_Hung")
HINHVE = os.path.join(REPORT, "Hinhve")
TEMPLATE = os.path.join(REPORT, "HUST_PPT_template_2022_blue_4x3.pptx")
OUT = os.path.join(REPORT, "DATN_VuManhHung_slides.pptx")
MD_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "slides_content.md")

HUST_BLUE = RGBColor(0x00, 0x3D, 0x7A)
ACCENT = RGBColor(0xC8, 0x10, 0x2E)
GREY = RGBColor(0x44, 0x44, 0x44)
GREEN = RGBColor(0x1B, 0x7A, 0x3D)
LIGHTBLUE = RGBColor(0xE7, 0xEF, 0xF7)
LIGHTGREEN = RGBColor(0xDD, 0xEC, 0xD9)

# name → fill cho các ô trong flow()
FLOW_FILLS = {"lightblue": LIGHTBLUE, "blue": LIGHTBLUE, "green": LIGHTGREEN}
# color name → (fill, text) cho @banner
BANNER_COLORS = {"green": (LIGHTGREEN, GREEN), "blue": (LIGHTBLUE, HUST_BLUE)}

# ---- Presentation base ----
COVER_SLIDE = None
_cover_sldid = None
_sldids_to_remove = []
if os.path.exists(TEMPLATE):
    prs = Presentation(TEMPLATE)
    USING_TEMPLATE = True
    KEEP_COVER = 2
    _all = list(prs.slides._sldIdLst)
    _cover_sldid = _all[KEEP_COVER]
    _sldids_to_remove = [sid for i, sid in enumerate(_all) if i != KEEP_COVER]
    COVER_SLIDE = prs.slides[KEEP_COVER]
else:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    USING_TEMPLATE = False

SW, SH = prs.slide_width, prs.slide_height
CONTENT_L = prs.slide_layouts[8] if USING_TEMPLATE else prs.slide_layouts[6]


def slide(layout=None):
    return prs.slides.add_slide(layout or CONTENT_L)


def clean_placeholders(s):
    for ph in list(s.placeholders):
        idx = ph.placeholder_format.idx
        typ = ph.placeholder_format.type
        if idx == 0 or typ == PP_PLACEHOLDER.SLIDE_NUMBER:
            continue
        ph._element.getparent().remove(ph._element)


def _tf(box):
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def set_lines(shape, lines, size=None, bold=None, color=None):
    tf = shape.text_frame
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    tf.paragraphs[0].clear()
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = ln
        if size is not None:
            r.font.size = Pt(size)
        if bold is not None:
            r.font.bold = bold
        if color is not None:
            r.font.color.rgb = color
    return shape


def title(s, text):
    if USING_TEMPLATE and s.shapes.title is not None:
        ph = s.shapes.title
        ph.text = text
        for p in ph.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            for r in p.runs:
                r.font.size = Pt(22); r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        clean_placeholders(s)
    else:
        box = s.shapes.add_textbox(Inches(0.5), Inches(0.3), SW - Inches(1.0), Inches(0.9))
        p = _tf(box).paragraphs[0]
        r = p.add_run(); r.text = text
        r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = HUST_BLUE
    return s


def _as_emu(v, default_emu):
    if v is None:
        return default_emu
    if isinstance(v, Length):
        return v
    return Inches(v)


def note(s, text):
    tf = s.notes_slide.notes_text_frame
    tf.text = text


def bullets(s, items, left=0.6, top=1.55, width=None, height=None, size=18, color=None):
    w = _as_emu(width, Inches(10 - 2 * left))
    h = _as_emu(height, Inches(5.4))
    box = s.shapes.add_textbox(Inches(left), Inches(top), w, h)
    tf = _tf(box)
    for i, it in enumerate(items):
        txt, lvl = (it if isinstance(it, tuple) else (it, 0))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = lvl
        r = p.add_run()
        r.text = ("• " if lvl == 0 else "– ") + txt
        r.font.size = Pt(size - lvl * 2)
        r.font.color.rgb = color or RGBColor(0x22, 0x22, 0x22)
        p.space_after = Pt(5)
    return s


def image(s, name, left, top, max_w, max_h, caption=None):
    path = os.path.join(HINHVE, name)
    if not os.path.exists(path):
        return placeholder(s, f"[thiếu ảnh: {name}]", left, top, max_w, max_h)
    pic = s.shapes.add_picture(path, Inches(left), Inches(top), width=Inches(max_w))
    if pic.height > Inches(max_h):
        ratio = Inches(max_h) / pic.height
        pic.width = int(pic.width * ratio); pic.height = Inches(max_h)
    pic.left = Inches(left) + int((Inches(max_w) - pic.width) / 2)
    if caption:
        cap = s.shapes.add_textbox(Inches(left), Inches(top) + pic.height + Pt(2), Inches(max_w), Inches(0.4))
        p = _tf(cap).paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = caption; r.font.size = Pt(11); r.font.italic = True; r.font.color.rgb = GREY
    return pic


def images_row(s, names, top=1.5, height=3.5, total_w=9.0, left0=0.5, gap=0.3, captions=None):
    n = len(names)
    w = (total_w - gap * (n - 1)) / n
    x = left0
    for i, nm in enumerate(names):
        cap = captions[i] if captions else None
        image(s, nm, x, top, w, height, caption=cap)
        x += w + gap


def _box(s, x, y, w, h, txt, fill, fs=12, tcolor=None):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = fill; b.line.color.rgb = HUST_BLUE
    tf = b.text_frame; tf.word_wrap = True; tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    for i, ln in enumerate(txt if isinstance(txt, list) else [txt]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = ln; r.font.size = Pt(fs if i == 0 else fs - 2)
        r.font.bold = (i == 0); r.font.color.rgb = tcolor or RGBColor(0x11, 0x11, 0x11)
    return b


def flow(s, steps, top, left=0.6, total_w=8.8, box_h=1.0, horizontal=True, fills=None, box_w=None):
    n = len(steps)
    lightblue = LIGHTBLUE
    if horizontal:
        gap = 0.5
        w = (total_w - gap * (n - 1)) / n
        x = left
        for i, txt in enumerate(steps):
            _box(s, x, top, w, box_h, txt, (fills[i] if fills else lightblue), fs=13)
            if i < n - 1:
                a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + w + 0.06),
                                       Inches(top + box_h / 2 - 0.12), Inches(gap - 0.12), Inches(0.24))
                a.fill.solid(); a.fill.fore_color.rgb = ACCENT; a.line.fill.background()
            x += w + gap
    else:
        gap = 0.35
        w = box_w or total_w
        cx = left + (total_w - w) / 2
        y = top
        for i, txt in enumerate(steps):
            _box(s, cx, y, w, box_h, txt, (fills[i] if fills else lightblue), fs=13)
            if i < n - 1:
                a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(left + total_w / 2 - 0.12),
                                       Inches(y + box_h + 0.02), Inches(0.24), Inches(gap - 0.06))
                a.fill.solid(); a.fill.fore_color.rgb = ACCENT; a.line.fill.background()
            y += box_h + gap


def placeholder(s, text, left, top, w, h):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor(0xEE, 0xF2, 0xF7); shp.line.color.rgb = HUST_BLUE
    p = _tf(shp).paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text; r.font.size = Pt(13); r.font.italic = True; r.font.color.rgb = HUST_BLUE
    return shp


def table(s, data, left, top, width, height, header=True, fs=13):
    rows, cols = len(data), len(data[0])
    gt = s.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(width), Inches(height)).table
    for r in range(rows):
        for c in range(cols):
            cell = gt.cell(r, c)
            cell.text = str(data[r][c])
            para = cell.text_frame.paragraphs[0]; para.font.size = Pt(fs)
            if header and r == 0:
                para.font.bold = True; para.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                cell.fill.solid(); cell.fill.fore_color.rgb = HUST_BLUE
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if r % 2 else RGBColor(0xEE, 0xF2, 0xF7)
    return gt


def img_slide(t, name, img_h=5.7, img_w=9.0, left=0.5, top=1.4, notes=None, caption=None):
    """Slide ảnh-lớn chiếm gần trọn, chữ giải thích để trong speaker notes."""
    s = slide(); title(s, t)
    image(s, name, left, top, img_w, img_h, caption=caption)
    if notes:
        note(s, notes)
    return s


def _heading(s, text, top=4.7, size=28):
    """Chữ lớn căn giữa (dùng cho slide Cảm ơn)."""
    box = s.shapes.add_textbox(Inches(0.8), Inches(top), SW - Inches(1.6), Inches(1.2))
    p = _tf(box).paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = HUST_BLUE
    return s


def divider(t, sub=None):
    """Slide phân khối (section divider): nền xanh HUST, chữ trắng căn giữa."""
    s = slide()
    clean_placeholders(s)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = HUST_BLUE; bg.line.fill.background()
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.2), Inches(4.15), Inches(3.6), Inches(0.06))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
    tb = s.shapes.add_textbox(Inches(0.6), Inches(2.7), SW - Inches(1.2), Inches(1.4))
    tf = _tf(tb)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = t
    r.font.size = Pt(34); r.font.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = sub
        r2.font.size = Pt(17); r2.font.color.rgb = RGBColor(0xCF, 0xDD, 0xEF)
    return s


# =================== PARSER MARKDOWN ===================
_ATTR_RE = re.compile(r'(\w+)=("[^"]*"|\S+)')


def parse_attrs(s):
    """Đọc 'key=val key="val có space"'. `caption=` đặc biệt: lấy hết phần còn lại của dòng."""
    s = s or ""
    caption = None
    if "caption=" in s:
        s, caption = s.split("caption=", 1)
        caption = caption.strip()
    d = {}
    for k, v in _ATTR_RE.findall(s):
        if v[:1] == '"':
            v = v[1:-1]
        else:
            low = v.lower()
            if low in ("true", "false"):
                v = (low == "true")
            else:
                try:
                    v = float(v)
                except ValueError:
                    pass
        d[k] = v
    if caption is not None:
        d["caption"] = caption
    return d


def parse_md(path):
    """Trả về list slide: {kind, title, sub, blocks:[{type, attrs, lines}]}."""
    slides, cur, block = [], None, None

    def close_block():
        nonlocal block
        if cur is not None and block is not None:
            cur["blocks"].append(block)
        block = None

    with open(path, encoding="utf-8") as f:
        for raw in f:
            line = raw.rstrip("\n")
            st = line.strip()
            if st.startswith("## "):
                close_block()
                if cur is not None:
                    slides.append(cur)
                parts = [p.strip() for p in st[3:].split("|")]
                cur = {
                    "kind": (parts[0].lower() if parts and parts[0] else "content"),
                    "title": parts[1] if len(parts) > 1 else "",
                    "sub": parts[2] if len(parts) > 2 else "",
                    "blocks": [],
                }
            elif st.startswith("### @"):
                close_block()
                head = st[4:]                       # bỏ '### '
                bits = head.split(None, 1)
                btype = bits[0][1:]                 # bỏ '@'
                block = {"type": btype, "attrs": parse_attrs(bits[1] if len(bits) > 1 else ""),
                         "lines": []}
            elif cur is not None and block is not None:
                block["lines"].append(line)
    close_block()
    if cur is not None:
        slides.append(cur)
    return slides


def read_bullets(lines):
    items = []
    for ln in lines:
        if not ln.strip():
            continue
        indent = len(ln) - len(ln.lstrip(" "))
        txt = ln.strip()
        if txt.startswith("- "):
            txt = txt[2:]
        lvl = 1 if indent >= 2 else 0
        items.append((txt, lvl) if lvl else txt)
    return items


def read_flow(lines):
    return [ln.strip()[2:].replace("\\n", "\n") for ln in lines if ln.strip().startswith("- ")]


def read_row(lines):
    names, caps = [], []
    for ln in lines:
        s = ln.strip()
        if not s.startswith("- "):
            continue
        body = s[2:]
        if "|" in body:
            n, c = body.split("|", 1)
            names.append(n.strip()); caps.append(c.strip())
        else:
            names.append(body.strip()); caps.append(None)
    return names, (caps if any(caps) else None)


def read_table(lines):
    data = []
    for ln in lines:
        s = ln.strip()
        if not s.startswith("|"):
            continue
        cells = [c.strip() for c in s.strip("|").split("|")]
        if cells and all(re.fullmatch(r"[-: ]+", c or "") for c in cells):
            continue  # dòng phân cách kiểu |---|
        data.append(cells)
    return data


def _joined(lines):
    return "\n".join(l for l in lines if l.strip()).replace("\\n", "\n")


# =================== RENDER ===================

def render_block(s, blk):
    t, a, lines = blk["type"], blk["attrs"], blk["lines"]
    if t == "bullets":
        bullets(s, read_bullets(lines),
                left=a.get("left", 0.6), top=a.get("top", 1.55),
                width=(Inches(a["width"]) if "width" in a else None),
                size=int(a.get("size", 18)))
    elif t == "image":
        image(s, a["name"], a.get("left", 0.5), a.get("top", 1.4),
              a.get("w", 9.0), a.get("h", 5.7), caption=a.get("caption"))
    elif t == "images-row":
        names, caps = read_row(lines)
        images_row(s, names, top=a.get("top", 1.5), height=a.get("height", 3.5),
                   total_w=a.get("total_w", 9.0), left0=a.get("left0", 0.5),
                   gap=a.get("gap", 0.3), captions=caps)
    elif t == "flow":
        fills = None
        if "fills" in a:
            fills = [FLOW_FILLS.get(x.strip(), LIGHTBLUE) for x in str(a["fills"]).split(",")]
        flow(s, read_flow(lines), top=a.get("top", 1.7), left=a.get("left", 0.6),
             total_w=a.get("total_w", 8.8), box_h=a.get("box_h", 1.0),
             horizontal=a.get("horizontal", True), fills=fills, box_w=a.get("box_w"))
    elif t == "table":
        gt = table(s, read_table(lines), a.get("left", 1.0), a.get("top", 1.5),
                   a.get("width", 8.0), a.get("height", 3.0), fs=int(a.get("fs", 13)))
        if "colw" in a:
            for i, w in enumerate(str(a["colw"]).split(",")):
                gt.columns[i].width = Inches(float(w))
    elif t == "banner":
        fill, tcol = BANNER_COLORS.get(a.get("color", "green"), (LIGHTGREEN, GREEN))
        _box(s, a.get("left", 0.6), a.get("top", 5.8), a.get("w", 8.8), a.get("h", 0.8),
             [l.strip() for l in lines if l.strip()], fill, fs=int(a.get("size", 16)), tcolor=tcol)
    elif t == "notes":
        note(s, _joined(lines))
    elif t == "heading":
        _heading(s, " ".join(l.strip() for l in lines if l.strip()),
                 top=a.get("top", 4.7), size=int(a.get("size", 28)))
    else:
        raise ValueError(f"Block @{t} chưa hỗ trợ")


def render_cover(sd):
    if not (USING_TEMPLATE and COVER_SLIDE is not None):
        return
    tlines, slines = [], []
    for blk in sd["blocks"]:
        if blk["type"] == "title":
            tlines = [l.strip() for l in blk["lines"] if l.strip()]
        elif blk["type"] == "subtitle":
            slines = [(l.strip()[2:] if l.strip().startswith("- ") else l.strip())
                      for l in blk["lines"] if l.strip()]
    for sh in COVER_SLIDE.shapes:
        if not sh.has_text_frame:
            continue
        u = sh.text_frame.text.upper()
        if "PRESENTATION TITLE" in u and tlines:
            set_lines(sh, tlines, size=24)
        elif "SUBTITLE" in u and slines:
            set_lines(sh, slines, size=15)
    clean_placeholders(COVER_SLIDE)


def render_md(path):
    for sd in parse_md(path):
        kind = sd["kind"]
        if kind == "cover":
            render_cover(sd)
        elif kind == "divider":
            divider(sd["title"], sd["sub"] or None)
        else:  # content
            s = slide()
            if sd["title"]:
                title(s, sd["title"])
            for blk in sd["blocks"]:
                render_block(s, blk)


# =================== BUILD ===================
render_md(MD_PATH)

# --- Dọn slide template thừa + đưa bìa lên đầu ---
if USING_TEMPLATE:
    lst = prs.slides._sldIdLst
    for sid in _sldids_to_remove:
        rId = sid.get(qn('r:id'))
        prs.part.drop_rel(rId)
        lst.remove(sid)
    lst.remove(_cover_sldid)
    lst.insert(0, _cover_sldid)

try:
    prs.save(OUT)
    print("SAVED:", OUT)
    print("Slides:", len(prs.slides._sldIdLst), "| Template:", "HUST" if USING_TEMPLATE else "default")
except PermissionError:
    print("PermissionError: đóng file DATN_VuManhHung_slides.pptx trước!")
